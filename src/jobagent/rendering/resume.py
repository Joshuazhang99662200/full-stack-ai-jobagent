"""Render a verified resume variant to PDF and Word, without changing a word of it.

Layout is decided here, content never is: every visible line is either a header
line the person supplied or the verbatim text of a verified item, in variant
order, so the rendered PDF passes `harness check-pdf`. The conventions are fixed:

- items whose id ends in `_ROLE` are the "company | title | dates" line of a role;
- items in a summary section are plain paragraphs;
- every other item is a bullet whose short lead before the first colon is bold.

reportlab and python-docx are optional; a missing one is reported, not guessed around.
"""

# ruff: noqa: RUF001  -- The Chinese full-width colon is a layout delimiter here.

from collections.abc import Collection, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Literal
from xml.sax.saxutils import escape

from jobagent.errors import ContractValidationError
from jobagent.schemas.harness import ResumeHeader
from jobagent.schemas.optimizer import ResumeVariant

DEFAULT_SUMMARY_SECTIONS = frozenset({"个人简介", "Summary"})
MAX_LEAD_CHARS = 14
NAVY = "1F3A5F"

# Regular / bold pairs tried in order when no font is given. CJK text needs one.
CJK_FONT_CANDIDATES: tuple[tuple[str, str], ...] = (
    (r"C:\Windows\Fonts\Deng.ttf", r"C:\Windows\Fonts\Dengb.ttf"),
    (r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\simhei.ttf"),
    (
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
    ),
    ("/System/Library/Fonts/STHeiti Medium.ttc", "/System/Library/Fonts/STHeiti Medium.ttc"),
)


@dataclass(frozen=True, slots=True)
class Block:
    kind: Literal["heading", "role", "paragraph", "bullet"]
    text: str
    lead: str | None = None


def _lead(text: str) -> str | None:
    for colon in ("：", ":"):
        index = text.find(colon)
        if 0 < index <= MAX_LEAD_CHARS:
            return text[: index + 1]
    return None


def layout(
    variant: ResumeVariant, *, summary_sections: Collection[str] = DEFAULT_SUMMARY_SECTIONS
) -> list[Block]:
    """Turn verified items into ordered blocks. No text is added or altered."""
    if not variant.verification.passed:
        raise ContractValidationError(
            "Only a verified resume variant can be rendered.",
            details={"variant": variant.id},
        )
    blocks: list[Block] = []
    section: str | None = None
    for item in variant.items:
        if item.section != section:
            section = item.section
            blocks.append(Block("heading", section))
        if item.id.endswith("_ROLE"):
            blocks.append(Block("role", item.text))
        elif section in summary_sections:
            blocks.append(Block("paragraph", item.text))
        else:
            blocks.append(Block("bullet", item.text, _lead(item.text)))
    return blocks


def default_cjk_font() -> tuple[str, str] | None:
    for regular, bold in CJK_FONT_CANDIDATES:
        if Path(regular).is_file() and Path(bold).is_file():
            return regular, bold
    return None


def render_pdf(
    variant: ResumeVariant,
    header: ResumeHeader,
    path: Path,
    *,
    font_path: str | None,
    bold_font_path: str | None = None,
    summary_sections: Collection[str] = DEFAULT_SUMMARY_SECTIONS,
) -> None:
    blocks = layout(variant, summary_sections=summary_sections)
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Table, TableStyle
    except ImportError as error:  # pragma: no cover - depends on the environment
        raise ContractValidationError(
            "PDF rendering needs reportlab; install the `render` extra."
        ) from error

    # Helvetica's "•" extracts as a control character; "·" survives text extraction.
    regular, bold, marker = "Helvetica", "Helvetica-Bold", "·"
    if font_path is not None:
        pdfmetrics.registerFont(TTFont("ResumeRegular", font_path))
        pdfmetrics.registerFont(TTFont("ResumeBold", bold_font_path or font_path))
        pdfmetrics.registerFontFamily("ResumeRegular", normal="ResumeRegular", bold="ResumeBold")
        regular, bold, marker = "ResumeRegular", "ResumeBold", "•"
    navy = colors.HexColor(f"#{NAVY}")
    name = ParagraphStyle("name", fontName=bold, fontSize=20, leading=25, textColor=navy)
    meta = ParagraphStyle("meta", fontName=regular, fontSize=10, leading=14.5, wordWrap="CJK")
    heading = ParagraphStyle(
        "heading", fontName=bold, fontSize=12.5, leading=16, textColor=navy, spaceBefore=7
    )
    role = ParagraphStyle("role", fontName=bold, fontSize=10.5, leading=15, spaceBefore=5)
    body = ParagraphStyle(
        "body", fontName=regular, fontSize=10, leading=14.5, wordWrap="CJK", spaceAfter=2
    )
    # The marker must use the registered font too, or reportlab falls back to Helvetica
    # and "•" extracts as a control character.
    bullet = ParagraphStyle(
        "bullet", parent=body, leftIndent=11, bulletIndent=1, bulletFontName=regular
    )

    story: list[object] = [Paragraph(escape(header.lines[0]), name)]
    for line in header.lines[1:]:
        left, _, right = line.partition("\t")
        if right:
            width = A4[0] - 29 * mm
            table = Table(
                [[Paragraph(escape(left), meta), Paragraph(escape(right), meta)]],
                colWidths=[width * 0.62, width * 0.38],
            )
            table.setStyle(
                TableStyle(
                    [
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("TOPPADDING", (0, 0), (-1, -1), 0),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                    ]
                )
            )
            story.append(table)
        else:
            story.append(Paragraph(escape(line), meta))
    for block in blocks:
        if block.kind == "heading":
            story += [
                Paragraph(escape(block.text), heading),
                HRFlowable(width="100%", thickness=0.9, color=navy, spaceBefore=1, spaceAfter=4),
            ]
        elif block.kind == "role":
            story.append(Paragraph(escape(block.text), role))
        elif block.kind == "paragraph":
            story.append(Paragraph(escape(block.text), body))
        else:
            text = escape(block.text)
            if block.lead is not None:
                rest = escape(block.text[len(block.lead) :])
                text = f"<b>{escape(block.lead)}</b>{rest}"
            story.append(Paragraph(text, bullet, bulletText=marker))
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=14 * mm,
        topMargin=13 * mm,
        bottomMargin=11 * mm,
        title=variant.id,
    ).build(story)


def render_docx(
    variant: ResumeVariant,
    header: ResumeHeader,
    path: Path,
    *,
    font_name: str,
    summary_sections: Collection[str] = DEFAULT_SUMMARY_SECTIONS,
) -> None:
    blocks = layout(variant, summary_sections=summary_sections)
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as error:  # pragma: no cover - depends on the environment
        raise ContractValidationError(
            "Word rendering needs python-docx; install the `render` extra."
        ) from error

    document = Document()
    for section in document.sections:
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
        section.top_margin, section.bottom_margin = Cm(1.3), Cm(1.1)
        section.left_margin, section.right_margin = Cm(1.5), Cm(1.4)
    style = document.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    navy = RGBColor.from_string(NAVY)

    def paragraph(style_name: str | None = None):  # type: ignore[no-untyped-def]
        created = document.add_paragraph(style=style_name)
        created.paragraph_format.space_after = Pt(2)
        created.paragraph_format.line_spacing = 1.05
        return created

    def run(target, text: str, *, bold: bool = False, size: float | None = None, color=None):  # type: ignore[no-untyped-def]
        added = target.add_run(text)
        added.bold = bold
        added.font.name = font_name
        added._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        if size is not None:
            added.font.size = Pt(size)
        if color is not None:
            added.font.color.rgb = color
        return added

    run(paragraph(), header.lines[0], bold=True, size=20, color=navy)
    for line in header.lines[1:]:
        left, tab, right = line.partition("\t")
        line_paragraph = paragraph()
        run(line_paragraph, left)
        if tab:
            line_paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(10))
            run(line_paragraph, "\t" + right)
    for block in blocks:
        if block.kind == "heading":
            created = paragraph()
            created.paragraph_format.space_before = Pt(7)
            run(created, block.text, bold=True, size=12.5, color=navy)
            _bottom_border(created, NAVY)
        elif block.kind == "role":
            created = paragraph()
            created.paragraph_format.space_before = Pt(4)
            run(created, block.text, bold=True, size=10.5)
        elif block.kind == "paragraph":
            created = paragraph()
            created.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run(created, block.text)
        else:
            created = paragraph("List Bullet")
            created.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if block.lead is not None:
                run(created, block.lead, bold=True)
                run(created, block.text[len(block.lead) :])
            else:
                run(created, block.text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _bottom_border(paragraph, color: str) -> None:  # type: ignore[no-untyped-def]
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(key), value)
    borders.append(bottom)
    properties.append(borders)


def block_texts(blocks: Sequence[Block]) -> list[str]:
    return [block.text for block in blocks]
