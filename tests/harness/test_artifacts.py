import zipfile
from xml.etree import ElementTree

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import Text
from reportlab.pdfgen import canvas

from jobagent.errors import ContractValidationError
from jobagent.harness.pdf import check_pdf
from tests.harness.test_pdf import render, variant


def test_pdf_rejects_internal_note_even_if_it_was_in_verified_text(tmp_path):
    document = variant()
    document.items[0].text = "Internal note: source needs owner confirmation."
    path = tmp_path / "resume.pdf"
    render(path, document.items[0].text)
    assert "internal_annotation" in check_pdf(document, path).issues


@pytest.mark.parametrize("extra", ["annotation", "attachment", "metadata"])
def test_pdf_rejects_non_employer_review_payload(tmp_path, extra):
    document = variant()
    path = tmp_path / "resume.pdf"
    render(path, document.items[0].text)
    writer = PdfWriter(clone_from=PdfReader(path))
    if extra == "annotation":
        writer.add_annotation(0, Text(rect=(10, 10, 100, 100), text="Reviewer comment"))
    elif extra == "attachment":
        writer.add_attachment("notes.txt", b"Private review notes")
    else:
        writer.add_metadata({"/Subject": "Internal note: private review context"})
    with path.open("wb") as output:
        writer.write(output)
    assert not check_pdf(document, path).passed


def test_pdf_preserves_normal_hyperlinks(tmp_path):
    document = variant()
    path = tmp_path / "resume.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 700, document.items[0].text)
    pdf.linkURL("https://example.com", (50, 695, 200, 710))
    pdf.save()
    assert check_pdf(document, path).passed


def test_internal_software_name_is_not_a_review_annotation(tmp_path):
    document = variant()
    document.items[0].text = "Built an internal note-taking tool for research teams."
    path = tmp_path / "resume.pdf"
    render(path, document.items[0].text)
    assert check_pdf(document, path).passed


def test_pdf_link_cannot_hide_javascript_action(tmp_path):
    from pypdf.generic import DictionaryObject, NameObject, TextStringObject

    document = variant()
    path = tmp_path / "resume.pdf"
    render(path, document.items[0].text)
    writer = PdfWriter(clone_from=PdfReader(path))
    annotation = {
        "/Type": "/Annot",
        "/Subtype": "/Link",
        "/Rect": [10, 10, 100, 100],
        "/A": DictionaryObject(
            {
                NameObject("/S"): NameObject("/JavaScript"),
                NameObject("/JS"): TextStringObject("console.println('synthetic');"),
            }
        ),
    }
    writer.add_annotation(0, annotation)
    with path.open("wb") as output:
        writer.write(output)
    assert "pdf_active_content" in check_pdf(document, path).issues


def test_pdf_xmp_metadata_cannot_carry_internal_notes(tmp_path):
    from pypdf.generic import DecodedStreamObject, NameObject

    document = variant()
    path = tmp_path / "resume.pdf"
    render(path, document.items[0].text)
    writer = PdfWriter(clone_from=PdfReader(path))
    metadata = DecodedStreamObject()
    metadata.set_data(b"<metadata>Internal note: private review context</metadata>")
    metadata.update({NameObject("/Type"): NameObject("/Metadata")})
    writer.root_object[NameObject("/Metadata")] = metadata
    with path.open("wb") as output:
        writer.write(output)
    assert "internal_annotation" in check_pdf(document, path).issues


def _docx(path, text):
    document = Document()
    document.add_paragraph(text)
    document.save(path)


@pytest.mark.parametrize("extra", ["comments", "revisions", "hidden", "metadata"])
def test_word_rejects_review_or_hidden_content(tmp_path, extra):
    from jobagent.harness.artifacts import check_artifact

    document = variant()
    path = tmp_path / "resume.docx"
    _docx(path, document.items[0].text)
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    if extra == "comments":
        members["word/comments.xml"] = (
            b'<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
        )
    elif extra == "metadata":
        metadata = ElementTree.fromstring(members["docProps/core.xml"])
        description = metadata.find("{http://purl.org/dc/elements/1.1/}description")
        assert description is not None
        description.text = "Internal note: draft only"
        members["docProps/core.xml"] = ElementTree.tostring(metadata)
    else:
        old = b"<w:p>"
        addition = (
            b'<w:ins w:id="1"><w:r><w:t>review</w:t></w:r></w:ins>'
            if extra == "revisions"
            else b"<w:r><w:rPr><w:vanish/></w:rPr><w:t>hidden</w:t></w:r>"
        )
        members["word/document.xml"] = members["word/document.xml"].replace(old, old + addition, 1)
    with zipfile.ZipFile(path, "w") as archive:
        for name, value in members.items():
            archive.writestr(name, value)
    report = check_artifact(path, variant=document)
    assert not report.passed


def test_clean_word_is_compared_with_verified_content(tmp_path):
    from jobagent.harness.artifacts import check_artifact

    document = variant()
    path = tmp_path / "resume.docx"
    _docx(path, document.items[0].text)
    assert check_artifact(path, variant=document).passed
    _docx(path, "An unapproved rewrite.")
    assert not check_artifact(path, variant=document).passed


def test_word_explicit_visible_formatting_is_permitted(tmp_path):
    from jobagent.harness.artifacts import check_artifact

    expected = variant()
    path = tmp_path / "resume.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run(expected.items[0].text).font.hidden = False
    document.save(path)
    assert check_artifact(path, variant=expected).passed


def test_actual_final_bytes_cannot_change_after_binding(tmp_path):
    from jobagent.harness.artifacts import bind_artifact, verify_artifact

    document = variant()
    path = tmp_path / "resume.pdf"
    render(path, document.items[0].text)
    binding = bind_artifact(path, document)
    verify_artifact(binding, document)
    path.write_bytes(path.read_bytes() + b"\n% changed since approval\n")
    with pytest.raises(ContractValidationError, match="changed"):
        verify_artifact(binding, document)


def test_binding_checks_do_not_modify_employer_file(tmp_path):
    from jobagent.harness.artifacts import check_artifact

    path = tmp_path / "resume.docx"
    _docx(path, "Internal note: remove this before sending.")
    before = path.read_bytes()
    assert not check_artifact(path).passed
    assert path.read_bytes() == before
