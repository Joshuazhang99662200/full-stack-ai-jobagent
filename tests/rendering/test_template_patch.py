import hashlib
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document


def source(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Candidate header")
    doc.add_paragraph("Original product narrative with a supported result.")
    doc.add_paragraph("Fixed following section")
    path = tmp_path / "source.docx"
    doc.save(path)
    return path


def test_word_patch_changes_only_authorized_text_and_bold(tmp_path: Path) -> None:
    from jobagent.rendering.template_patch import patch_template
    from jobagent.schemas.template_tailoring import ParagraphEdit, TemplatePatch, WordRun

    original = source(tmp_path)
    before = original.read_bytes()
    plan = TemplatePatch(
        template_sha256=hashlib.sha256(before).hexdigest(),
        edits=[ParagraphEdit(
            paragraph_index=1,
            expected_text="Original product narrative with a supported result.",
            runs=[WordRun(text="Product outcome: ", bold=True), WordRun(text="supported result.")],
        )],
    )
    output = tmp_path / "tailored.docx"
    patch_template(original, output, plan, allowed_paragraphs={1})
    assert original.read_bytes() == before
    with ZipFile(original) as a, ZipFile(output) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name != "word/document.xml":
                assert a.read(name) == b.read(name), name
    doc = Document(output)
    assert [p.text for p in doc.paragraphs] == [
        "Candidate header", "Product outcome: supported result.", "Fixed following section",
    ]
    assert doc.paragraphs[1].runs[0].bold is True
    assert doc.paragraphs[1].runs[1].bold is False


@pytest.mark.parametrize("problem", ["digest", "paragraph", "original", "overflow"])
def test_word_patch_refuses_changed_source_or_out_of_scope_edits(
    tmp_path: Path, problem: str,
) -> None:
    from jobagent.rendering.template_patch import patch_template
    from jobagent.schemas.template_tailoring import ParagraphEdit, TemplatePatch, WordRun

    original = source(tmp_path)
    plan = TemplatePatch(
        template_sha256=("0" * 64 if problem == "digest"
                         else hashlib.sha256(original.read_bytes()).hexdigest()),
        edits=[ParagraphEdit(
            paragraph_index=0 if problem == "paragraph" else 1,
            expected_text="wrong" if problem == "original" else
            "Original product narrative with a supported result.",
            runs=[WordRun(text="x" * 200 if problem == "overflow" else "New narrative")],
        )],
    )
    with pytest.raises(ValueError):
        patch_template(original, tmp_path / "out.docx", plan, allowed_paragraphs={1})
    assert not (tmp_path / "out.docx").exists()


def test_patch_refuses_to_flatten_mixed_non_bold_styles(tmp_path: Path) -> None:
    from jobagent.rendering.template_patch import patch_template
    from jobagent.schemas.template_tailoring import ParagraphEdit, TemplatePatch, WordRun

    original = source(tmp_path)
    document = Document(original)
    document.paragraphs[1].add_run("Styled suffix").italic = True
    document.save(original)
    plan = TemplatePatch(
        template_sha256=hashlib.sha256(original.read_bytes()).hexdigest(),
        edits=[ParagraphEdit(paragraph_index=1, expected_text=document.paragraphs[1].text,
                             runs=[WordRun(text="Shorter result")])],
    )
    with pytest.raises(ValueError, match="non-bold"):
        patch_template(original, tmp_path / "out.docx", plan, allowed_paragraphs={1})


def test_explicit_run_edits_preserve_mixed_properties(tmp_path: Path) -> None:
    from jobagent.rendering.template_patch import patch_template
    from jobagent.schemas.template_tailoring import ParagraphEdit, RunTextEdit, TemplatePatch

    original = source(tmp_path)
    document = Document(original)
    document.paragraphs[1].add_run("Styled suffix").italic = True
    document.save(original)
    plan = TemplatePatch(
        template_sha256=hashlib.sha256(original.read_bytes()).hexdigest(),
        edits=[ParagraphEdit(paragraph_index=1, expected_text=document.paragraphs[1].text,
                             run_edits=[RunTextEdit(run_index=1, expected_text="Styled suffix",
                                                    text="New suffix")])],
    )
    output = tmp_path / "out.docx"
    patch_template(original, output, plan, allowed_paragraphs={1})
    result = Document(output)
    assert result.paragraphs[1].runs[0].text == document.paragraphs[1].runs[0].text
    assert result.paragraphs[1].runs[1].italic is True
    assert result.paragraphs[1].runs[1].text == "New suffix"
